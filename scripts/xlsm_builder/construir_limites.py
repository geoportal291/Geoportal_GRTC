# -*- coding: utf-8 -*-
"""
Construye la plantilla base (xlsx, pre-VBA) del Informe Excel de
Límites de Consistencia (MTC E 110 / E 111 / E 108), replicando el
formato del informe PDF del Geoportal GRTC.

La plantilla lleva:
  - Reporte     : maqueta del informe con fórmulas vivas
  - _DATOS      : hoja oculta con datos auxiliares de los gráficos
  - _MAPA       : hoja oculta ruta->celda que lee el generador Python
  - _GRAFICOS   : hoja oculta con la especificación de gráficos (macro VBA)

Salida: backend/templates/base/base_limites.xlsx
(el .xlsm final se crea con inyectar_vba.ps1)
"""
import os
import sys
sys.path.insert(0, os.path.dirname(__file__))

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.axis import ChartLines
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from comun import (RAIZ, DIR_PLANTILLAS, OSCURO, SECCION, CALC_FILL,
                   RESULT_FILL, BLANCO, F_LAB, F_TAB_HEAD, F_TAB_CELDA,
                   F_TAB_CALC, F_RESULT_LAB, F_RESULT_VAL, F_FIRMA,
                   AL_CENTRO, banda, celda, par_ficha, definir_nombre,
                   hoja_oculta_mapa, hoja_specs_graficos,
                   cabecera_institucional, preparar_hoja)

SALIDA = os.path.join(DIR_PLANTILLAS, "base", "base_limites.xlsx")

FMT_PESO = "0.00"
FMT_PCT = "0.00"
FMT_FECHA = 'dddd d" de "mmmm" de "yyyy'

IP_DESCRIPCION = [
    ("0 - 3", "No plástico"),
    ("3 - 15", "Ligeramente plástico"),
    ("15 - 30", "Baja plasticidad"),
    ("> 30", "Alta plasticidad"),
]


def tabla_transpuesta(ws, fila, params, n_cols, col_ini=2, nombre_cols=None):
    """
    Dibuja una tabla transpuesta (filas=parámetros, columnas=muestras).
    params: lista de (etiqueta, tipo, formato) tipo: input|calc_texto|calc_num
    Devuelve {indice_param: fila}.
    """
    celda(ws, fila, 1, "Parámetro", font=F_TAB_HEAD, fill=OSCURO)
    for j, nombre in enumerate(nombre_cols or [str(i + 1) for i in range(n_cols)]):
        celda(ws, fila, col_ini + j, nombre, font=F_TAB_HEAD, fill=OSCURO)
    ws.row_dimensions[fila].height = 16
    filas = {}
    for i, (etiqueta, tipo, formato) in enumerate(params):
        f = fila + 1 + i
        celda(ws, f, 1, etiqueta)
        for j in range(n_cols):
            cel = celda(ws, f, col_ini + j)
            if tipo == "calc_num":
                cel.fill = PatternFill("solid", fgColor=CALC_FILL)
                cel.font = F_TAB_CALC
            if formato:
                cel.number_format = formato
        filas[i] = f
    return filas


def construir():
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    preparar_hoja(ws, [34] + [11] * 8)  # A + B..I

    mapa = []  # (ruta, celda, tipo, formato)

    # ------------------------------------------------ cabecera institucional
    cabecera_institucional(ws, ultima_col=9)
    ws.row_dimensions[7].height = 4

    # ------------------------------------------------ franja de código
    fila = 8
    banda(ws, fila, 1, 1, "CÓDIGO", fill=OSCURO)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=9)
    for c in range(2, 10):
        celda(ws, fila, c, fill=OSCURO)
    cod = ws.cell(row=fila, column=2)
    cod.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cod.alignment = AL_CENTRO
    mapa.append(("meta.codigo_ensayo", f"B{fila}", "texto", None))
    ws.row_dimensions[fila].height = 20
    ws.row_dimensions[9].height = 6

    # ------------------------------------------------ ficha técnica
    fila = 10
    par_ficha(ws, fila, 1, "PROYECTO", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=9)
    for c in range(2, 10):
        celda(ws, fila, c, fill=RESULT_FILL)
    mapa.append(("meta.proyecto_nombre", f"B{fila}", "texto", None))

    fila = 11
    par_ficha(ws, fila, 1, "UBICACIÓN")
    par_ficha(ws, fila, 2, "Lugar:")
    par_ficha(ws, fila, 4, "Distrito:")
    par_ficha(ws, fila, 6, "Provincia:")
    par_ficha(ws, fila, 8, "Dpto:")
    mapa += [("meta.tramo_nombre", f"C{fila}", "texto", None),
             ("meta.distrito", f"E{fila}", "texto", None),
             ("meta.provincia", f"G{fila}", "texto", None),
             ("meta.departamento", f"I{fila}", "texto", None)]

    fila = 12
    par_ficha(ws, fila, 1, "SOLICITANTE", ancho_val=3)
    ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=6)
    celda(ws, fila, 5, "Fecha Muestreo:", font=F_LAB, fill="F8FAFC")
    celda(ws, fila, 6, fill="F8FAFC")
    ws.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=9)
    celda(ws, fila, 7, font=Font(name="Calibri", size=9, bold=True))
    mapa += [("meta.solicitante", f"B{fila}", "texto", None),
             ("meta.fecha_muestreo", f"G{fila}", "fecha", FMT_FECHA)]

    fila = 13
    par_ficha(ws, fila, 1, "COORDENADAS")
    par_ficha(ws, fila, 2, "E:")
    ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
    par_ficha(ws, fila, 4, None)
    par_ficha(ws, fila, 5, "N:")
    ws.merge_cells(start_row=fila, start_column=6, end_row=fila, end_column=7)
    par_ficha(ws, fila, 7, None)
    par_ficha(ws, fila, 8, "Prof. (m):")
    mapa += [("meta.coordenada_este", f"C{fila}", "numero", "0.00"),
             ("meta.coordenada_norte", f"F{fila}", "numero", "0.00"),
             ("meta.profundidad", f"I{fila}", "texto", None)]

    fila = 14
    par_ficha(ws, fila, 1, "DATOS DE MUESTRA")
    par_ficha(ws, fila, 2, "Explor.:")
    par_ficha(ws, fila, 4, "Progresiva:")
    par_ficha(ws, fila, 6, "Estrato:")
    par_ficha(ws, fila, 8, "Lado:")
    mapa += [("meta.calicata", f"C{fila}", "texto", None),
             ("meta.progresiva", f"E{fila}", "texto", None),
             ("meta.estrato", f"G{fila}", "texto", None),
             ("meta.lado", f"I{fila}", "texto", None)]
    for _f in range(10, 15):            # ficha técnica compacta y pareja
        ws.row_dimensions[_f].height = 20
    ws.row_dimensions[15].height = 6

    # ------------------------------------ sección Límite Líquido (MTC E 110)
    fila = 16
    banda(ws, fila, 1, 9, "DETERMINACIÓN DEL LÍMITE LÍQUIDO - MTC E 110")
    f_head = fila + 1
    params_ll = [
        ("Cod. Recipiente", "input", None),
        ("N° de Golpes", "input", "0"),
        ("Peso del recipiente (g)", "input", FMT_PESO),
        ("Recip. + suelo húmedo (g)", "input", FMT_PESO),
        ("Recip. + suelo seco (g)", "input", FMT_PESO),
        ("Peso agua (g)", "calc_num", FMT_PESO),
        ("Peso suelo seco (g)", "calc_num", FMT_PESO),
        ("% de Humedad", "calc_num", FMT_PCT),
        ("LL Corregido (%)", "calc_num", FMT_PCT),
    ]
    filas_ll = tabla_transpuesta(ws, f_head, params_ll, 3)
    f_cod, f_golpes = filas_ll[0], filas_ll[1]
    f_tara, f_humedo, f_seco = filas_ll[2], filas_ll[3], filas_ll[4]
    f_agua, f_ssec, f_hum, f_llc = filas_ll[5], filas_ll[6], filas_ll[7], filas_ll[8]
    for j, col in enumerate(("B", "C", "D")):
        n = j + 1
        mapa.append((f"datos.tables.limite_liquido.{n}.codigo", f"{col}{f_cod}", "texto", None))
        mapa.append((f"datos.tables.limite_liquido.{n}.golpes", f"{col}{f_golpes}", "numero", "0"))
        mapa.append((f"datos.tables.limite_liquido.{n}.tara", f"{col}{f_tara}", "numero", FMT_PESO))
        mapa.append((f"datos.tables.limite_liquido.{n}.humedo", f"{col}{f_humedo}", "numero", FMT_PESO))
        mapa.append((f"datos.tables.limite_liquido.{n}.seco", f"{col}{f_seco}", "numero", FMT_PESO))
        ws[f"{col}{f_agua}"] = f"=IF(COUNT({col}{f_humedo}:{col}{f_seco})=2,{col}{f_humedo}-{col}{f_seco},\"\")"
        ws[f"{col}{f_ssec}"] = f"=IF(COUNT({col}{f_tara},{col}{f_seco})=2,{col}{f_seco}-{col}{f_tara},\"\")"
        ws[f"{col}{f_hum}"] = (f"=IF(AND(ISNUMBER({col}{f_ssec}),N({col}{f_ssec})>0),"
                               f"IFERROR({col}{f_agua}/{col}{f_ssec}*100,0),\"\")")
        ws[f"{col}{f_llc}"] = (f"=IF(AND(ISNUMBER({col}{f_hum}),N({col}{f_golpes})>0),"
                               f"{col}{f_hum}*POWER({col}{f_golpes}/25,0.121),\"\")")

    f_ll_res = filas_ll[8] + 1
    par_ficha(ws, f_ll_res, 1, "LÍMITE LÍQUIDO (L.L.)", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=f_ll_res, start_column=2, end_row=f_ll_res, end_column=5)
    for c in range(2, 6):
        celda(ws, f_ll_res, c, fill=RESULT_FILL)
    cel = ws.cell(row=f_ll_res, column=2)
    cel.value = (f"=IF(COUNTIF(B{f_golpes}:D{f_golpes},25)>0,"
                 f"INDEX(B{f_hum}:D{f_hum},MATCH(25,B{f_golpes}:D{f_golpes},0)),"
                 f"IF(COUNT(B{f_llc}:D{f_llc})>0,AVERAGE(B{f_llc}:D{f_llc}),0))")
    cel.font = F_RESULT_VAL
    cel.number_format = '0.00" %"'
    cel.alignment = AL_CENTRO
    f_ll_res_celda = f"B{f_ll_res}"
    ws.row_dimensions[f_ll_res + 1].height = 6

    # ------------------------------------ sección Límite Plástico (MTC E 111)
    fila = f_ll_res + 2
    banda(ws, fila, 1, 9, "DETERMINACIÓN DEL LÍMITE PLÁSTICO - MTC E 111")
    params_lp = [
        ("Cod. Recipiente", "input", None),
        ("Peso del recipiente (g)", "input", FMT_PESO),
        ("Recip. + suelo húmedo (g)", "input", FMT_PESO),
        ("Recip. + suelo seco (g)", "input", FMT_PESO),
        ("Peso agua (g)", "calc_num", FMT_PESO),
        ("Peso suelo seco (g)", "calc_num", FMT_PESO),
        ("% de Humedad", "calc_num", FMT_PCT),
    ]
    filas_lp = tabla_transpuesta(ws, fila + 1, params_lp, 2)
    f_cod2, f_tara2 = filas_lp[0], filas_lp[1]
    f_humedo2, f_seco2 = filas_lp[2], filas_lp[3]
    f_agua2, f_ssec2, f_hum2 = filas_lp[4], filas_lp[5], filas_lp[6]
    for j, col in enumerate(("B", "C")):
        n = j + 1
        mapa += [
            (f"datos.tables.limite_plastico.{n}.codigo", f"{col}{f_cod2}", "texto", None),
            (f"datos.tables.limite_plastico.{n}.tara", f"{col}{f_tara2}", "numero", FMT_PESO),
            (f"datos.tables.limite_plastico.{n}.humedo", f"{col}{f_humedo2}", "numero", FMT_PESO),
            (f"datos.tables.limite_plastico.{n}.seco", f"{col}{f_seco2}", "numero", FMT_PESO),
        ]
        ws[f"{col}{f_agua2}"] = f"=IF(COUNT({col}{f_humedo2}:{col}{f_seco2})=2,{col}{f_humedo2}-{col}{f_seco2},\"\")"
        ws[f"{col}{f_ssec2}"] = f"=IF(COUNT({col}{f_tara2},{col}{f_seco2})=2,{col}{f_seco2}-{col}{f_tara2},\"\")"
        ws[f"{col}{f_hum2}"] = (f"=IF(AND(ISNUMBER({col}{f_ssec2}),N({col}{f_ssec2})>0),"
                                f"IFERROR({col}{f_agua2}/{col}{f_ssec2}*100,0),\"\")")

    f_lp_res = f_hum2 + 1
    par_ficha(ws, f_lp_res, 1, "LÍMITE PLÁSTICO (L.P.)", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=f_lp_res, start_column=2, end_row=f_lp_res, end_column=5)
    for c in range(2, 6):
        celda(ws, f_lp_res, c, fill=RESULT_FILL)
    cel = ws.cell(row=f_lp_res, column=2)
    cel.value = f"=IF(COUNT(B{f_hum2}:C{f_hum2})>0,AVERAGE(B{f_hum2}:C{f_hum2}),0)"
    cel.font = F_RESULT_VAL
    cel.number_format = '0.00" %"'
    cel.alignment = AL_CENTRO
    f_lp_res_celda = f"B{f_lp_res}"

    f_ip_res = f_lp_res + 1
    par_ficha(ws, f_ip_res, 1, "ÍNDICE DE PLASTICIDAD (I.P.)", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=f_ip_res, start_column=2, end_row=f_ip_res, end_column=5)
    for c in range(2, 6):
        celda(ws, f_ip_res, c, fill=RESULT_FILL)
    cel = ws.cell(row=f_ip_res, column=2)
    cel.value = f"=IF(AND(ISNUMBER({f_ll_res_celda}),ISNUMBER({f_lp_res_celda})),{f_ll_res_celda}-{f_lp_res_celda},\"\")"
    cel.font = F_RESULT_VAL
    cel.number_format = '0.00" %"'
    cel.alignment = AL_CENTRO
    f_ip_res_celda = f"B{f_ip_res}"

    # tabla estática IP / Descripción
    f_ip_tab = f_ip_res + 1
    ws.merge_cells(start_row=f_ip_tab, start_column=1, end_row=f_ip_tab, end_column=2)
    celda(ws, f_ip_tab, 1, "IP", font=F_TAB_HEAD, fill=OSCURO)
    celda(ws, f_ip_tab, 2, fill=OSCURO)
    ws.merge_cells(start_row=f_ip_tab, start_column=3, end_row=f_ip_tab, end_column=5)
    celda(ws, f_ip_tab, 3, "Descripción", font=F_TAB_HEAD, fill=OSCURO)
    for c in (4, 5):
        celda(ws, f_ip_tab, c, fill=OSCURO)
    for i, (rango, desc) in enumerate(IP_DESCRIPCION):
        f = f_ip_tab + 1 + i
        ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=2)
        celda(ws, f, 1, rango)
        celda(ws, f, 2)
        ws.merge_cells(start_row=f, start_column=3, end_row=f, end_column=5)
        celda(ws, f, 3, desc)
        for c in (4, 5):
            celda(ws, f, c)
    ws.row_dimensions[f_ip_tab + len(IP_DESCRIPCION) + 1].height = 6

    # ------------------------------------ sección Humedad Natural (MTC E 108)
    fila = f_ip_tab + len(IP_DESCRIPCION) + 2
    banda(ws, fila, 1, 9, "ENSAYO DE HUMEDAD NATURAL - MTC E 108")
    params_hn = [
        ("Cod. Cápsula", "input", None),
        ("Peso de cápsula (g)", "input", FMT_PESO),
        ("Cáps. + suelo húmedo (g)", "input", FMT_PESO),
        ("Cáps. + suelo seco (g)", "input", FMT_PESO),
        ("Peso del agua (g)", "calc_num", FMT_PESO),
        ("Peso del suelo seco (g)", "calc_num", FMT_PESO),
        ("Contenido de humedad (%)", "calc_num", FMT_PCT),
    ]
    filas_hn = tabla_transpuesta(ws, fila + 1, params_hn, 2,
                                 nombre_cols=["Muestra 1", "Muestra 2"])
    f_cod3, f_tara3 = filas_hn[0], filas_hn[1]
    f_humedo3, f_seco3 = filas_hn[2], filas_hn[3]
    f_agua3, f_ssec3, f_hum3 = filas_hn[4], filas_hn[5], filas_hn[6]
    for j, col in enumerate(("B", "C")):
        n = j + 1
        mapa += [
            (f"datos.tables.humedad_natural.{n}.codigo", f"{col}{f_cod3}", "texto", None),
            (f"datos.tables.humedad_natural.{n}.tara", f"{col}{f_tara3}", "numero", FMT_PESO),
            (f"datos.tables.humedad_natural.{n}.humedo", f"{col}{f_humedo3}", "numero", FMT_PESO),
            (f"datos.tables.humedad_natural.{n}.seco", f"{col}{f_seco3}", "numero", FMT_PESO),
        ]
        ws[f"{col}{f_agua3}"] = f"=IF(COUNT({col}{f_humedo3}:{col}{f_seco3})=2,{col}{f_humedo3}-{col}{f_seco3},\"\")"
        ws[f"{col}{f_ssec3}"] = f"=IF(COUNT({col}{f_tara3},{col}{f_seco3})=2,{col}{f_seco3}-{col}{f_tara3},\"\")"
        ws[f"{col}{f_hum3}"] = (f"=IF(AND(ISNUMBER({col}{f_ssec3}),N({col}{f_ssec3})>0),"
                                f"IFERROR({col}{f_agua3}/{col}{f_ssec3}*100,0),\"\")")

    f_hn_res = f_hum3 + 1
    par_ficha(ws, f_hn_res, 1, "HUMEDAD NATURAL (w %)", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=f_hn_res, start_column=2, end_row=f_hn_res, end_column=5)
    for c in range(2, 6):
        celda(ws, f_hn_res, c, fill=RESULT_FILL)
    cel = ws.cell(row=f_hn_res, column=2)
    cel.value = f"=IF(COUNT(B{f_hum3}:C{f_hum3})>0,AVERAGE(B{f_hum3}:C{f_hum3}),0)"
    cel.font = F_RESULT_VAL
    cel.number_format = '0.00" %"'
    cel.alignment = AL_CENTRO
    f_hn_res_celda = f_hn_res
    ws.row_dimensions[f_hn_res + 1].height = 6

    # ------------------------------------------------------------- gráficos
    f_graf = f_hn_res + 2
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=f_graf - 1))
    banda(ws, f_graf, 1, 9, "GRÁFICOS DEL ENSAYO")

    # --------------------------------- hoja _DATOS (auxiliar de gráficos)
    hd = wb.create_sheet("_DATOS")
    hd.sheet_state = "hidden"
    hd["F1"] = "=IFERROR(SLOPE(B3:B5,C3:C5),0)"
    hd["F2"] = "=IFERROR(INTERCEPT(B3:B5,C3:C5),0)"
    hd["F3"] = "=F1*LOG10(25)+F2"
    hd["F4"] = f"=Reporte!{f_ll_res_celda}"
    hd["F5"] = f"=Reporte!{f_lp_res_celda}"
    hd["F6"] = "=F4-F5"
    hd["H3"] = 25
    hd["I3"] = "=F3"
    hd["H5"], hd["I5"] = 20, 0
    hd["H6"], hd["I6"] = 100, 58.4
    xs = [10, 12.5, 15, 20, 25, 35, 50]
    for i, x in enumerate(xs):
        hd.cell(row=3 + i, column=4, value=x)
        hd.cell(row=3 + i, column=5, value=f"=$F$1*LOG10(D{3 + i})+$F$2")
    for i, col in enumerate(("B", "C", "D")):
        r = 3 + i
        hd[f"A{r}"] = f"=IF(OR(Reporte!{col}{f_golpes}=\"\",N(Reporte!{col}{f_golpes})=0),NA(),Reporte!{col}{f_golpes})"
        hd[f"B{r}"] = f"=IF(ISNA(A{r}),NA(),Reporte!{col}{f_hum})"

        hd[f"C{r}"] = f"=IF(ISNA(A{r}),NA(),LOG10(A{r}))"
    nombres = {
        "FluidezGolpesX": "$A$3:$A$5", "FluidezHumedadY": "$B$3:$B$5",
        "FluidezLogX": "$C$3:$C$5", "RegresionX": "$D$3:$D$9",
        "RegresionY": "$E$3:$E$9", "PuntoLLX": "$H$3", "PuntoLLY": "$I$3",
        "CartaLLX": "$F$4", "CartaIPY": "$F$6",
        "LineaAX": "$H$5:$H$6", "LineaAY": "$I$5:$I$6",
    }
    for nombre, ref in nombres.items():
        definir_nombre(wb, nombre, "_DATOS", ref)

    # ------------------------------------ gráfico Curva de Fluidez (openpyxl)
    ch = ScatterChart()
    ch.title = "Curva de Fluidez — % Humedad vs N° de Golpes (esc. log)"
    ch.style = 2
    ch.height = 11.2
    ch.width = 12.0
    ch.x_axis.scaling.logBase = 10
    ch.x_axis.scaling.min = 10
    ch.x_axis.scaling.max = 50
    ch.x_axis.majorGridlines = ChartLines()
    ch.y_axis.majorGridlines = ChartLines()
    ch.legend.position = "b"
    ch.legend.overlay = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.varyColors = False

    def serie(chart, x_ref, y_ref, titulo, color, linea=False, guion=False,
              marker=False, size=7):
        xv = Reference(hd, range_string=f"_DATOS!{x_ref}")
        yv = Reference(hd, range_string=f"_DATOS!{y_ref}")
        s = Series(yv, xv, title=titulo)
        lp = LineProperties(solidFill=color, w=22000)
        if guion:
            lp.prstDash = "dash"
        if linea:
            s.graphicalProperties.line = lp
        else:
            s.graphicalProperties.line = LineProperties(noFill=True)
        if marker:
            s.marker = Marker(symbol="circle", size=size)
            s.marker.graphicalProperties.solidFill = color
            s.marker.graphicalProperties.line = LineProperties(solidFill=color)
        else:
            s.marker = Marker(symbol="none")
        s.smooth = False
        chart.series.append(s)

    serie(ch, "$D$3:$D$9", "$E$3:$E$9", "Línea de Fluidez (Regresión)",
          "DC2626", linea=True, guion=True)
    serie(ch, "$A$3:$A$5", "$B$3:$B$5", "Puntos de Ensayo", "1E40AF", marker=True)
    serie(ch, "$H$3:$H$3", "$I$3:$I$3", "L.L. (25 golpes)", "059669", marker=True, size=9)
    ws.add_chart(ch, f"A{f_graf + 1}")

    # --------------------------------- gráfico Carta de Plasticidad (openpyxl)
    ch2 = ScatterChart()
    ch2.title = "Carta de Plasticidad — IP vs Límite Líquido"
    ch2.style = 2
    ch2.height = 11.2
    ch2.width = 8.3
    ch2.x_axis.scaling.min = 0
    ch2.x_axis.scaling.max = 100
    ch2.y_axis.scaling.min = 0
    ch2.y_axis.scaling.max = 60
    ch2.x_axis.majorGridlines = ChartLines()
    ch2.y_axis.majorGridlines = ChartLines()
    ch2.legend.position = "b"
    ch2.legend.overlay = False
    ch2.varyColors = False
    ch2.x_axis.delete = False
    ch2.y_axis.delete = False
    serie(ch2, "$H$5:$H$6", "$I$5:$I$6", "Línea A", "475569", linea=True, guion=True)
    serie(ch2, "$F$4:$F$4", "$F$6:$F$6", "Muestra", "7C3AED", marker=True, size=8)
    ws.add_chart(ch2, f"F{f_graf + 1}")

    # ---------------------------------------------- especificación para VBA
    RG = f"'_DATOS'!"
    series_specs = [
        ("curva_fluidez", "Determinación del Límite Líquido (Curva de Fluidez)",
         "xy", f"A{f_graf + 1}", 455, 425, "SI", 10, 50, "", "",
         "Número de Golpes", "% de Humedad",
         "Línea de Fluidez (Regresión)", "linea", f"{RG}$D$3:$D$9", f"{RG}$E$3:$E$9", "DC2626", 1),
        ("curva_fluidez", "", "", "", "", "", "", "", "", "", "", "", "",
         "Puntos de Ensayo", "marcador", f"{RG}$A$3:$A$5", f"{RG}$B$3:$B$5", "1E40AF", 2),
        ("curva_fluidez", "", "", "", "", "", "", "", "", "", "", "", "",
         "L.L. (25 golpes)", "marcador", f"{RG}$H$3:$H$3", f"{RG}$I$3:$I$3", "059669", 3),
        ("carta_plasticidad", "Carta de Plasticidad de Casagrande",
         "xy", f"F{f_graf + 1}", 315, 425, "NO", 0, 100, 0, 60,
         "Límite Líquido (LL)", "Índice de Plasticidad (IP)",
         "Línea A", "linea", f"{RG}$H$5:$H$6", f"{RG}$I$5:$I$6", "475569", 1),
        ("carta_plasticidad", "", "", "", "", "", "", "", "", "", "", "", "",
         "Muestra", "marcador", f"{RG}$F$4:$F$4", f"{RG}$F$6:$F$6", "7C3AED", 2),
    ]
    hoja_specs_graficos(wb, series_specs)

    # ------------------------------------------------------------- firmas
    f_firma = f_graf + 30
    bloques = [("ESP. ENSAYOS GEOTECNICOS", 1, 3),
               ("ESP. SUELOS Y PAVIMENTOS", 4, 6),
               ("SUPERVISOR", 7, 9)]
    for _, ci, cf in bloques:
        for c in range(ci, cf + 1):
            ws.cell(row=f_firma, column=c).border = Border(
                top=Side(style="thin", color="0F172A"))
    for texto, ci, cf in bloques:
        ws.merge_cells(start_row=f_firma + 1, start_column=ci, end_row=f_firma + 1, end_column=cf)
        cel = ws.cell(row=f_firma + 1, column=ci, value=texto)
        cel.font = F_FIRMA
        cel.alignment = AL_CENTRO

    ws.print_area = f"A1:I{f_firma + 2}"

    # ------------------------------------------------------------- _MAPA
    hoja_oculta_mapa(wb, mapa)

    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    print(f"OK -> {SALIDA}  (filas: informe hasta {f_firma + 2}, mapa: {len(mapa)} rutas)")


if __name__ == "__main__":
    construir()
