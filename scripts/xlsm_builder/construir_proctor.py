# -*- coding: utf-8 -*-
"""
Construye la plantilla base (xlsx, pre-VBA) del Informe Excel de
Proctor Modificado (MTC E 115) — mismo estilo que el informe PDF.

Salida: backend/templates/base/base_proctor.xlsx
"""
import os
import sys
sys.path.insert(0, os.path.dirname(__file__))

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.axis import ChartLines
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Border, Font, PatternFill, Side

from comun import (DIR_PLANTILLAS, OSCURO, CALC_FILL, RESULT_FILL,
                   F_LAB, F_TAB_HEAD, F_TAB_CALC, F_RESULT_VAL, F_FIRMA,
                   AL_CENTRO, banda, celda, par_ficha, definir_nombre,
                   hoja_oculta_mapa, hoja_specs_graficos,
                   cabecera_institucional, preparar_hoja)

SALIDA = os.path.join(DIR_PLANTILLAS, "base", "base_proctor.xlsx")
FMT_PESO = "0.00"
FMT_DENS = "0.000"
FMT_FECHA = 'dddd d" de "mmmm" de "yyyy'

MUESTRAS = ["M1", "M2", "M3", "M4"]


def construir():
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    preparar_hoja(ws, [34] + [11] * 8)

    mapa = []

    # ------------------------------------- cabecera, código y ficha técnica
    cabecera_institucional(ws, ultima_col=9)
    ws.row_dimensions[7].height = 4
    fila = 8
    banda(ws, fila, 1, 1, "CÓDIGO", fill=OSCURO)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=9)
    for c in range(2, 10):
        celda(ws, fila, c, fill=OSCURO)
    cod = ws.cell(row=fila, column=2)
    cod.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cod.alignment = AL_CENTRO
    mapa.append(("meta.codigo_ensayo", f"B{fila}", "texto", None))
    ws.row_dimensions[fila].height = 20
    ws.row_dimensions[9].height = 6

    fila = 10
    par_ficha(ws, fila, 1, "PROYECTO", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=9)
    for c in range(2, 10):
        celda(ws, fila, c, fill=RESULT_FILL)
    mapa.append(("meta.proyecto_nombre", f"B{fila}", "texto", None))
    fila = 11
    par_ficha(ws, fila, 1, "UBICACIÓN")
    par_ficha(ws, fila, 2, "Lugar:")
    par_ficha(ws, fila, 4, "Distrito:")
    par_ficha(ws, fila, 6, "Provincia:")
    par_ficha(ws, fila, 8, "Dpto:")
    mapa += [("meta.tramo_nombre", f"C{fila}", "texto", None),
             ("meta.distrito", f"E{fila}", "texto", None),
             ("meta.provincia", f"G{fila}", "texto", None),
             ("meta.departamento", f"I{fila}", "texto", None)]
    fila = 12
    par_ficha(ws, fila, 1, "SOLICITANTE", ancho_val=3)
    ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=6)
    celda(ws, fila, 5, "Fecha Muestreo:", font=F_LAB, fill="F8FAFC")
    celda(ws, fila, 6, fill="F8FAFC")
    ws.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=9)
    celda(ws, fila, 7, font=Font(name="Calibri", size=9, bold=True))
    mapa += [("meta.solicitante", f"B{fila}", "texto", None),
             ("meta.fecha_muestreo", f"G{fila}", "fecha", FMT_FECHA)]
    fila = 13
    par_ficha(ws, fila, 1, "COORDENADAS")
    par_ficha(ws, fila, 2, "E:")
    ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
    par_ficha(ws, fila, 5, "N:")
    ws.merge_cells(start_row=fila, start_column=6, end_row=fila, end_column=7)
    par_ficha(ws, fila, 8, "Prof. (m):")
    mapa += [("meta.coordenada_este", f"C{fila}", "numero", "0.00"),
             ("meta.coordenada_norte", f"F{fila}", "numero", "0.00"),
             ("meta.profundidad", f"I{fila}", "texto", None)]
    fila = 14
    par_ficha(ws, fila, 1, "DATOS DE MUESTRA")
    par_ficha(ws, fila, 2, "Explor.:")
    par_ficha(ws, fila, 4, "Progresiva:")
    par_ficha(ws, fila, 6, "Estrato:")
    par_ficha(ws, fila, 8, "Lado:")
    mapa += [("meta.calicata", f"C{fila}", "texto", None),
             ("meta.progresiva", f"E{fila}", "texto", None),
             ("meta.estrato", f"G{fila}", "texto", None),
             ("meta.lado", f"I{fila}", "texto", None)]
    for _f in range(10, 15):            # ficha técnica compacta y pareja
        ws.row_dimensions[_f].height = 20
    # valores de ubicación/coordenadas a 9pt: sin saltos de línea que
    # descuadren la ficha ("La Convención", "0.00 - 0.20 m")
    for _ref in ("C11", "E11", "G11", "I11", "C13", "F13", "I13"):
        ws[_ref].font = Font(name="Calibri", size=9, bold=True)
    ws.row_dimensions[15].height = 6

    # ---------------------------------------------- equipo y compactación
    fila = 16
    banda(ws, fila, 1, 9, "COMPACTACIÓN PROCTOR MODIFICADO - MTC E 115")
    fila = 17
    par_ficha(ws, fila, 1, "EQUIPO / MOLDE")
    par_ficha(ws, fila, 2, "Código:")
    par_ficha(ws, fila, 4, "Diám. (cm):")
    par_ficha(ws, fila, 6, "Altura (cm):")
    par_ficha(ws, fila, 8, "Vol. (cm³):")
    ws[f"I{fila}"] = '=IF(AND(N(E17)>0,N(G17)>0),ROUND(PI()*POWER(E17,2)*G17/4,1),"")'
    ws[f"I{fila}"].fill = PatternFill("solid", fgColor=CALC_FILL)
    ws[f"I{fila}"].font = F_TAB_CALC
    mapa += [
        ("datos.general_fields.codigo_molde", f"C{fila}", "texto", None),
        ("datos.general_fields.diametro_cm", f"E{fila}", "numero", FMT_PESO),
        ("datos.general_fields.altura_cm", f"G{fila}", "numero", FMT_PESO),
    ]
    fila = 18
    par_ficha(ws, fila, 1, "COMPACTACIÓN")
    par_ficha(ws, fila, 2, "Golpes/capa:")
    par_ficha(ws, fila, 4, "N° capas:")
    par_ficha(ws, fila, 6, "Martillo (lb):")
    mapa += [
        ("datos.general_fields.golpes_por_capa", f"C{fila}", "numero", "0"),
        ("datos.general_fields.numero_capas", f"E{fila}", "numero", "0"),
        ("datos.general_fields.masa_martillo_lb", f"G{fila}", "numero", "0"),
    ]
    ws.row_dimensions[19].height = 6

    # ----------------------------------------------- tabla datos de ensayo
    # cada muestra ocupa un par de columnas para llegar al ancho A..I
    PARES = [(2, 3), (4, 5), (6, 7), (8, 9)]   # B:C, D:E, F:G, H:I
    LETRAS = ["B", "D", "F", "H"]               # celda ancla del valor

    def tabla_dos(ws, fila_h, titulo, params, calc_fill=True):
        banda(ws, fila_h - 1, 1, 9, titulo)
        celda(ws, fila_h, 1, "Parámetro", font=F_TAB_HEAD, fill=OSCURO)
        for j, (c1, c2) in enumerate(PARES):
            ws.merge_cells(start_row=fila_h, start_column=c1, end_row=fila_h, end_column=c2)
            celda(ws, fila_h, c1, MUESTRAS[j], font=F_TAB_HEAD, fill=OSCURO)
            celda(ws, fila_h, c2, fill=OSCURO)
        ws.row_dimensions[fila_h].height = 16
        filas = {}
        for i, (etiqueta, es_calc, formato) in enumerate(params):
            f = fila_h + 1 + i
            celda(ws, f, 1, etiqueta)
            for (c1, c2) in PARES:
                ws.merge_cells(start_row=f, start_column=c1, end_row=f, end_column=c2)
                cel = celda(ws, f, c1)
                celda(ws, f, c2)
                if es_calc and calc_fill:
                    cel.fill = PatternFill("solid", fgColor=CALC_FILL)
                    ws.cell(row=f, column=c2).fill = PatternFill("solid", fgColor=CALC_FILL)
                    cel.font = F_TAB_CALC
                if formato:
                    cel.number_format = formato
            filas[i] = f
        return filas

    f_datos = tabla_dos(ws, 21, "DATOS DE ENSAYO", [
        ("Masa del molde (g)", False, FMT_PESO),
        ("Masa molde + suelo húmedo (g)", False, FMT_PESO),
        ("Agua incrementada (g)", False, FMT_PESO),
        ("Masa del suelo húmedo (g)", True, FMT_PESO),
        ("Volumen del molde (cm³)", True, FMT_PESO),
        ("Densidad húmeda (g/cm³)", True, FMT_DENS),
    ])
    f_molde, f_molde_suelo, f_agua_inc = f_datos[0], f_datos[1], f_datos[2]
    f_suelo_hum, f_vol, f_dens_hum = f_datos[3], f_datos[4], f_datos[5]
    for j, col in enumerate(LETRAS):
        n = j + 1
        mapa += [
            (f"datos.tables.datos_ensayo.m{n}.masa_molde", f"{col}{f_molde}", "numero", FMT_PESO),
            (f"datos.tables.datos_ensayo.m{n}.masa_suelo_humedo_molde", f"{col}{f_molde_suelo}", "numero", FMT_PESO),
            (f"datos.tables.datos_ensayo.m{n}.agua_incrementada", f"{col}{f_agua_inc}", "numero", FMT_PESO),
        ]
        ws[f"{col}{f_suelo_hum}"] = f"=IF(COUNT({col}{f_molde}:{col}{f_molde_suelo})=2,{col}{f_molde_suelo}-{col}{f_molde},\"\")"
        ws[f"{col}{f_vol}"] = '=IF(N($I$17)>0,$I$17,2124)'
        ws[f"{col}{f_dens_hum}"] = (f"=IF(AND(ISNUMBER({col}{f_suelo_hum}),N({col}{f_vol})>0),"
                                    f"IFERROR({col}{f_suelo_hum}/{col}{f_vol},0),\"\")")
    ws.row_dimensions[f_datos[5] + 1].height = 6

    # ------------------------------------------ tabla cálculo de la humedad
    f_hum_tbl = f_dens_hum + 3
    f_h = tabla_dos(ws, f_hum_tbl + 1, "CÁLCULO DE LA HUMEDAD", [
        ("Cápsula N°", False, None),
        ("Masa de cápsula (g)", False, FMT_PESO),
        ("Masa cápsula + suelo húmedo (g)", False, FMT_PESO),
        ("Masa cápsula + suelo seco (g)", False, FMT_PESO),
        ("Masa de agua (g)", True, FMT_PESO),
        ("Masa de suelo seco (g)", True, FMT_PESO),
        ("Contenido de humedad (%)", True, FMT_PESO),
        ("Densidad seca (g/cm³)", True, FMT_DENS),
    ])
    f_caps, f_mcap = f_h[0], f_h[1]
    f_mcap_hum, f_mcap_sec = f_h[2], f_h[3]
    f_magua, f_msec, f_humedad, f_dens_sec = f_h[4], f_h[5], f_h[6], f_h[7]
    for j, col in enumerate(LETRAS):
        n = j + 1
        mapa += [
            (f"datos.tables.calculo_humedad.m{n}.capsula_nro", f"{col}{f_caps}", "texto", None),
            (f"datos.tables.calculo_humedad.m{n}.masa_capsula", f"{col}{f_mcap}", "numero", FMT_PESO),
            (f"datos.tables.calculo_humedad.m{n}.masa_capsula_suelo_humedo", f"{col}{f_mcap_hum}", "numero", FMT_PESO),
            (f"datos.tables.calculo_humedad.m{n}.masa_capsula_suelo_seco", f"{col}{f_mcap_sec}", "numero", FMT_PESO),
        ]
        ws[f"{col}{f_magua}"] = f"=IF(COUNT({col}{f_mcap_hum}:{col}{f_mcap_sec})=2,{col}{f_mcap_hum}-{col}{f_mcap_sec},\"\")"
        ws[f"{col}{f_msec}"] = f"=IF(COUNT({col}{f_mcap},{col}{f_mcap_sec})=2,{col}{f_mcap_sec}-{col}{f_mcap},\"\")"
        ws[f"{col}{f_humedad}"] = (f"=IF(AND(ISNUMBER({col}{f_msec}),N({col}{f_msec})>0),"
                                   f"IFERROR({col}{f_magua}/{col}{f_msec}*100,0),\"\")")
        ws[f"{col}{f_dens_sec}"] = (f"=IF(AND(ISNUMBER({col}{f_humedad}),ISNUMBER({col}{f_dens_hum})),"
                                    f"IFERROR({col}{f_dens_hum}/(1+{col}{f_humedad}/100),0),\"\")")

    # ------------------------------------------------------- resultados MDS/OCH
    f_mds = f_dens_sec + 2
    par_ficha(ws, f_mds, 1, "MÁXIMA DENSIDAD SECA (g/cm³)", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=f_mds, start_column=2, end_row=f_mds, end_column=9)
    for c in range(2, 10):
        celda(ws, f_mds, c, fill=RESULT_FILL)
    cel = ws.cell(row=f_mds, column=2)
    cel.value = "=IF(ISNUMBER(_DATOS!$F$5),_DATOS!$F$5,\"\")"
    cel.font = F_RESULT_VAL
    cel.number_format = '0.000" g/cm³"'
    cel.alignment = AL_CENTRO

    f_och = f_mds + 1
    par_ficha(ws, f_och, 1, "HUMEDAD ÓPTIMA (%)", fill_val=RESULT_FILL)
    ws.merge_cells(start_row=f_och, start_column=2, end_row=f_och, end_column=9)
    for c in range(2, 10):
        celda(ws, f_och, c, fill=RESULT_FILL)
    cel = ws.cell(row=f_och, column=2)
    cel.value = "=IF(ISNUMBER(_DATOS!$F$4),_DATOS!$F$4,\"\")"
    cel.font = F_RESULT_VAL
    cel.number_format = '0.00" %"'
    cel.alignment = AL_CENTRO
    ws.row_dimensions[f_och + 1].height = 6

    # ---------------------------------------------------------------- gráfico
    f_graf = f_och + 3
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=f_graf - 1))
    banda(ws, f_graf, 1, 9, "CURVA DE COMPACTACIÓN")

    hd = wb.create_sheet("_DATOS")
    hd.sheet_state = "hidden"
    # puntos de ensayo (x = humedad, y = densidad seca)
    for j, col in enumerate(LETRAS):
        r = 3 + j
        hd[f"A{r}"] = (f"=IF(OR(NOT(ISNUMBER(Reporte!{col}{f_humedad})),"
                       f"NOT(ISNUMBER(Reporte!{col}{f_dens_sec})),Reporte!{col}{f_dens_sec}=0),"
                       f"NA(),Reporte!{col}{f_humedad})")
        hd[f"B{r}"] = (f"=IF(ISNA(A{r}),NA(),Reporte!{col}{f_dens_sec})")
    # ajuste cuadrático y = a x² + b x + c (vértice = MDS/OCH)
    hd["F1"] = '=IFERROR(INDEX(LINEST(B3:B6,A3:A6^{1,2}),1,1),"")'
    hd["F2"] = '=IFERROR(INDEX(LINEST(B3:B6,A3:A6^{1,2}),1,2),"")'
    hd["F3"] = '=IFERROR(INDEX(LINEST(B3:B6,A3:A6^{1,2}),1,3),"")'
    hd["F4"] = '=IF(AND(ISNUMBER(F1),N(F1)<>0),-F2/(2*F1),"")'
    hd["F5"] = '=IF(ISNUMBER(F4),F1*F4^2+F2*F4+F3,"")'
    for i in range(20):  # puntos de la parábola alrededor del óptimo
        r = 3 + i
        hd[f"H{r}"] = f'=IF(OR(NOT(ISNUMBER($F$1)),NOT(ISNUMBER($F$4))),NA(),$F$4+(ROW()-13)*0.5)'
        hd[f"I{r}"] = f'=IF(ISNA(H{r}),NA(),$F$1*H{r}^2+$F$2*H{r}+$F$3)'
    hd["K3"] = "=IF(ISNUMBER($F$4),$F$4,NA())"
    hd["L3"] = "=IF(ISNUMBER($F$5),$F$5,NA())"
    for nombre, ref in {"CurvaX": "$A$3:$A$6", "CurvaY": "$B$3:$B$6",
                        "AjusteX": "$H$3:$H$22", "AjusteY": "$I$3:$I$22",
                        "MDSX": "$K$3", "MDSY": "$L$3"}.items():
        definir_nombre(wb, nombre, "_DATOS", ref)

    ch = ScatterChart()
    ch.title = "Curva de Compactación — Densidad Seca (g/cm³) vs Humedad (%)"
    ch.height = 12.6
    ch.width = 16.4
    ch.x_axis.majorGridlines = ChartLines()
    ch.y_axis.majorGridlines = ChartLines()
    ch.legend.position = "b"
    ch.legend.overlay = False
    ch.varyColors = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False

    def serie(chart, x_ref, y_ref, titulo, color, linea=False, suave=False,
              marker=False, size=7):
        xv = Reference(hd, range_string=f"_DATOS!{x_ref}")
        yv = Reference(hd, range_string=f"_DATOS!{y_ref}")
        s = Series(yv, xv, title=titulo)
        if linea or suave:
            s.graphicalProperties.line = LineProperties(solidFill=color, w=24000)
        else:
            s.graphicalProperties.line = LineProperties(noFill=True)
        if marker:
            s.marker = Marker(symbol="circle", size=size)
            s.marker.graphicalProperties.solidFill = color
            s.marker.graphicalProperties.line = LineProperties(solidFill=color)
        else:
            s.marker = Marker(symbol="none")
        s.smooth = bool(suave)
        chart.series.append(s)

    serie(ch, "$H$3:$H$22", "$I$3:$I$22", "Ajuste Cuadrático", "DC2626", suave=True)
    serie(ch, "$A$3:$A$6", "$B$3:$B$6", "Puntos de Ensayo", "1E40AF", marker=True, size=8)
    serie(ch, "$K$3:$K$3", "$L$3:$L$3", "Máxima Densidad Seca", "059669", marker=True, size=10)
    ws.add_chart(ch, f"A{f_graf + 1}")

    RG = "'_DATOS'!"
    series_specs = [
        ("curva_compactacion", "Curva de Compactación Proctor Modificado",
         "xy", f"A{f_graf + 1}", 620, 480, "NO", "", "", "", "",
         "Contenido de Humedad (%)", "Densidad Seca (g/cm³)",
         "Ajuste Cuadrático", "linea_suave", f"{RG}$H$3:$H$22", f"{RG}$I$3:$I$22", "DC2626", 1),
        ("curva_compactacion", "", "", "", "", "", "", "", "", "", "", "", "",
         "Puntos de Ensayo", "marcador", f"{RG}$A$3:$A$6", f"{RG}$B$3:$B$6", "1E40AF", 2),
        ("curva_compactacion", "", "", "", "", "", "", "", "", "", "", "", "",
         "Máxima Densidad Seca", "marcador", f"{RG}$K$3:$K$3", f"{RG}$L$3:$L$3", "059669", 3),
    ]
    hoja_specs_graficos(wb, series_specs)

    # ---------------------------------------------------------------- firmas
    f_firma = f_graf + 31
    bloques = [("ESP. ENSAYOS GEOTECNICOS", 1, 3),
               ("ESP. SUELOS Y PAVIMENTOS", 4, 6),
               ("SUPERVISOR", 7, 9)]
    for _, ci, cf in bloques:
        for c in range(ci, cf + 1):
            ws.cell(row=f_firma, column=c).border = Border(
                top=Side(style="thin", color="0F172A"))
    for texto, ci, cf in bloques:
        ws.merge_cells(start_row=f_firma + 1, start_column=ci, end_row=f_firma + 1, end_column=cf)
        cel = ws.cell(row=f_firma + 1, column=ci, value=texto)
        cel.font = F_FIRMA
        cel.alignment = AL_CENTRO

    ws.print_area = f"A1:I{f_firma + 2}"
    hoja_oculta_mapa(wb, mapa)
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    print(f"OK -> {SALIDA}  (filas: hasta {f_firma + 2}, mapa: {len(mapa)} rutas)")


if __name__ == "__main__":
    construir()
