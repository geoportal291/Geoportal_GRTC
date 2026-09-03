# -*- coding: utf-8 -*-
"""
Utilidades compartidas para construir las plantillas .xlsm de informes
de ensayos de mecánica de suelos (Geoportal GRTC).

Las plantillas se generan con openpyxl (formato + fórmulas + gráficos) y
luego scripts/xlsm_builder/inyectar_vba.ps1 les incrusta las macros VBA
guardándolas como .xlsm.
"""
import os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

# ---------------------------------------------------------------- constantes
RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DIR_PLANTILLAS = os.path.join(RAIZ, "backend", "templates")
DIR_LOGOS = os.path.join(RAIZ, "frontend", "public", "assets")
LOGO_GRTC = os.path.join(DIR_LOGOS, "logo_grtc.png")
LOGO_CUSCO = os.path.join(DIR_LOGOS, "logo_cusco.png")

# Paleta del informe PDF
OSCURO = "1F2937"        # banda de código / cabeceras de tabla
SECCION = "334155"       # títulos de sección
GRIS_LABEL = "64748B"
CALC_FILL = "EEF2F7"     # filas calculadas
RESULT_FILL = "F1F5F9"   # cajas de resultado
FICHA_FILL = "F8FAFC"
BLANCO = "FFFFFF"
VERDE_OK = "059669"

BORDE_FINO = Side(style="thin", color="0F172A")
BORDE_MED = Side(style="medium", color="0F172A")
BORDE_TABLA = Border(left=BORDE_FINO, right=BORDE_FINO, top=BORDE_FINO, bottom=BORDE_FINO)

F_TITULO = Font(name="Calibri", size=14, bold=True, color="111827")
F_SUBTITULO = Font(name="Calibri", size=9, color="1F2937")
F_MINI = Font(name="Calibri", size=8, color="1F2937")
F_LEMA = Font(name="Calibri", size=7.5, italic=True, color="374151")
F_LAB = Font(name="Calibri", size=8, bold=True, color="475569")
F_VALOR = Font(name="Calibri", size=10, bold=True, color="111827")
F_SECCION = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_TAB_HEAD = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
F_TAB_CELDA = Font(name="Calibri", size=9, color="111827")
F_TAB_CALC = Font(name="Calibri", size=9, color="1D4ED8")
F_RESULT_LAB = Font(name="Calibri", size=9, bold=True, color="334155")
F_RESULT_VAL = Font(name="Calibri", size=11, bold=True, color="111827")
F_FIRMA = Font(name="Calibri", size=8, bold=True, color="334155")

AL_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_IZQ = Alignment(horizontal="left", vertical="center")
AL_DER = Alignment(horizontal="right", vertical="center")


def banda(ws, fila, col_ini, col_fin, texto, fill=SECCION, font=F_SECCION, alto=None):
    """Fila-banda con fondo de color y texto (títulos de sección, código)."""
    ws.merge_cells(start_row=fila, start_column=col_ini, end_row=fila, end_column=col_fin)
    for c in range(col_ini, col_fin + 1):
        cel = ws.cell(row=fila, column=c)
        cel.fill = PatternFill("solid", fgColor=fill)
        cel.border = BORDE_TABLA
    cel = ws.cell(row=fila, column=col_ini, value=texto)
    cel.font = font
    cel.alignment = AL_CENTRO
    if alto:
        ws.row_dimensions[fila].height = alto


def celda(ws, fila, col, valor=None, font=F_TAB_CELDA, fill=None, borde=True,
          align=None, formato=None):
    cel = ws.cell(row=fila, column=col)
    if valor is not None:
        cel.value = valor
    cel.font = font
    if fill:
        cel.fill = PatternFill("solid", fgColor=fill)
    if borde:
        cel.border = BORDE_TABLA
    cel.alignment = align or AL_CENTRO
    if formato:
        cel.number_format = formato
    return cel


def par_ficha(ws, fila, col, etiqueta, valor=None, ancho_val=1, formato=None,
              editable=False, fill_val=BLANCO):
    """Par etiqueta:valor en línea (ficha técnica). La etiqueta ocupa 1 col."""
    celda(ws, fila, col, etiqueta, font=F_LAB, fill=FICHA_FILL, align=AL_DER)
    if ancho_val > 1:
        ws.merge_cells(start_row=fila, start_column=col + 1,
                       end_row=fila, end_column=col + ancho_val)
        for cc in range(col + 1, col + ancho_val + 1):
            celda(ws, fila, cc, fill=fill_val)
    cel = ws.cell(row=fila, column=col + 1)
    if valor is not None:
        cel.value = valor
    cel.font = F_VALOR
    cel.alignment = AL_CENTRO
    if formato:
        cel.number_format = formato
    return cel


def definir_nombre(wb, nombre, hoja, ref):
    """Crea un rango con nombre (usado por gráficos y macros)."""
    if nombre in wb.defined_names:
        del wb.defined_names[nombre]
    wb.defined_names.add(DefinedName(nombre, attr_text=f"'{hoja}'!{ref}"))


def hoja_oculta_mapa(wb, filas, nombre_hoja="_MAPA"):
    """
    Hoja oculta con el mapa ruta->celda que lee el generador Python.
    filas: lista de (ruta, celda, tipo, formato)
    """
    ws = wb.create_sheet(nombre_hoja)
    ws.append(["ruta", "celda", "tipo", "formato"])
    for f in filas:
        ws.append(list(f))
    ws.sheet_state = "hidden"
    ws.sheet_properties.tabColor = "94A3B8"
    return ws


def hoja_specs_graficos(wb, series, nombre_hoja="_GRAFICOS"):
    """
    Hoja oculta con la especificación de gráficos que lee la macro VBA
    (y el builder openpyxl para crearlos). Una fila por serie:
    (id, titulo, tipo, ancla, ancho_px, alto_px, x_log, x_min, x_max,
     y_min, y_max, titulo_x, titulo_y, serie, estilo_serie, rango_x,
     rango_y, color_hex, orden)
    estilo_serie: linea | linea_suave | marcador
    """
    ws = wb.create_sheet(nombre_hoja)
    cab = ["id", "titulo", "tipo", "ancla", "ancho_px", "alto_px", "x_log",
           "x_min", "x_max", "y_min", "y_max", "titulo_x", "titulo_y",
           "serie", "estilo", "rango_x", "rango_y", "color", "orden"]
    ws.append(cab)
    for s in series:
        ws.append(list(s))
    ws.sheet_state = "hidden"
    ws.sheet_properties.tabColor = "94A3B8"
    return ws


def cabecera_institucional(ws, ultima_col):
    """Encabezado GRTC con logos y lema (filas 1-6)."""
    col_letra_fin = get_column_letter(ultima_col)
    col_medio = get_column_letter(max(3, (ultima_col - 1) // 2))
    lineas = [
        ("GOBIERNO REGIONAL CUSCO", F_TITULO),
        ("Gerencia Regional de Transportes y Comunicaciones Cusco", F_SUBTITULO),
        ("SUB GERENCIA DE COBERTURA Y COMUNICACIONES", F_MINI),
        ("UNIDAD FUNCIONAL ESTUDIOS Y PROYECTOS", F_MINI),
        ("Laboratorio de Mecánica de Suelos, Materiales y Pavimentos", F_SUBTITULO),
        ("«Año de la recuperación y consolidación de la economía peruana»", F_LEMA),
    ]
    for i, (texto, fuente) in enumerate(lineas, start=1):
        ws.merge_cells(f"C{i}:{col_letra_fin if ultima_col <= 7 else 'F'}{i}")
        cel = ws.cell(row=i, column=3, value=texto)
        cel.font = fuente
        cel.alignment = AL_CENTRO
    for fila_logo, ruta in ((1, LOGO_GRTC), (1, LOGO_CUSCO)):
        if os.path.exists(ruta):
            from openpyxl.drawing.image import Image as XImage
            img = XImage(ruta)
            escala = 52.0 / max(img.height, 1)
            img.height = int(img.height * escala)
            img.width = int(img.width * escala)
            img.anchor = "A1" if ruta == LOGO_GRTC else f"{get_column_letter(ultima_col - 1)}1"
            ws.add_image(img)
    for f in range(1, 7):
        ws.row_dimensions[f].height = 15


def preparar_hoja(ws, anchos, orient="portrait"):
    """Anchuras de columnas, configuración de página A4."""
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.page_setup.orientation = orient
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.sheet_view.showGridLines = False
