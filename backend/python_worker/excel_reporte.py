#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de informes Excel (.xlsm) con macros para ensayos de mecánica
de suelos (Geoportal GRTC).

Uso:
  python excel_reporte.py --plantilla <plantilla.xlsm> --payload <payload.json> --salida <salida.xlsm>

El payload JSON tiene la forma:
  {
    "meta":  { "codigo_ensayo": ..., "proyecto_nombre": ..., ... },
    "datos": { ... contenido de datos_formulario del ensayo ... }
  }

La plantilla contiene la hoja oculta _MAPA con la relación ruta->celda
(columnas: ruta, celda, tipo, formato). Los cálculos NO se hacen aquí:
son fórmulas vivas de la plantilla que Excel recalcula al abrir
(fullCalcOnLoad) y la macro VBA integrada regenera los gráficos.
"""
import argparse
import datetime
import json
import sys

from openpyxl import load_workbook


def resolver(dic, ruta):
    actual = dic
    for parte in ruta.split("."):
        if isinstance(actual, dict) and parte in actual:
            actual = actual[parte]
        else:
            return None
    return actual


def a_valor(val, tipo):
    if tipo == "numero":
        try:
            return float(val)
        except (TypeError, ValueError):
            return val
    if tipo == "fecha":
        if isinstance(val, datetime.datetime):
            return val
        if isinstance(val, datetime.date):
            return datetime.datetime(val.year, val.month, val.day)
        texto = str(val)[:10]
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.datetime.strptime(texto, fmt)
            except ValueError:
                continue
        return val
    return val


def main():
    ap = argparse.ArgumentParser(description="Genera informe .xlsm desde plantilla + payload")
    ap.add_argument("--plantilla", required=True)
    ap.add_argument("--payload", required=True)
    ap.add_argument("--salida", required=True)
    args = ap.parse_args()

    with open(args.payload, encoding="utf-8") as f:
        payload = json.load(f)

    try:
        wb = load_workbook(args.plantilla, keep_vba=True)
    except Exception as e:
        print(f"ERROR abriendo plantilla: {e}", file=sys.stderr)
        sys.exit(2)

    if "_MAPA" not in wb.sheetnames:
        print("ERROR: la plantilla no tiene hoja _MAPA", file=sys.stderr)
        sys.exit(2)
    if "Reporte" not in wb.sheetnames:
        print("ERROR: la plantilla no tiene hoja Reporte", file=sys.stderr)
        sys.exit(2)

    ws = wb["Reporte"]
    escritos = 0
    for ruta, celda_ref, tipo, formato in wb["_MAPA"].iter_rows(min_row=2, values_only=True):
        if not ruta or not celda_ref:
            continue
        raiz_p, _, resto = ruta.partition(".")
        fuente = payload.get(raiz_p)
        if fuente is None:
            continue
        val = resolver(fuente, resto) if resto else fuente
        if val is None or val == "":
            continue
        celda = ws[celda_ref]
        celda.value = a_valor(val, tipo or "texto")
        if formato:
            celda.number_format = formato
        escritos += 1

    wb.calculation.fullCalcOnLoad = True
    wb.save(args.salida)
    print(f"OK celdas={escritos}")


if __name__ == "__main__":
    main()
