# -*- coding: utf-8 -*-
"""
Genera la Ficha de Calicata en Excel (equivalente a la hoja "Res" del
formato MTC): consolidado de resultados de un estrato con encabezado
institucional GRTC, resumen de ensayos, descripción geotécnica y firmas.

Uso:
    python ficha_calicata_excel.py --payload ficha.json --salida ficha.xlsx

El payload lo arma backend/services/fichaCalicataService.js. Todo valor
ausente se imprime como "—": la ficha se genera aunque falten ensayos.
openpyxl puro, sin VBA: el Excel es solo el producto de impresión.
"""
import argparse
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ------------------------------------------------------------------ estilo
OSCURO = "1F2937"
SECCION = "334155"
RESULT_FILL = "E8EDF3"
CALC_FILL = "EEF2F7"
LABEL_FILL = "F5F7F9"
LABEL_TXT = "1C3D5F"
BORDE_TABLA = Border(*[Side(style="thin", color="CBD5E1")] * 4)

F_SECCION = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_LAB = Font(name="Calibri", size=9, bold=True, color=LABEL_TXT)
F_VAL = Font(name="Calibri", size=10, bold=True, color="0F172A")
F_CELDA = Font(name="Calibri", size=9)
F_CALC = Font(name="Calibri", size=9, bold=True, color="1E3A5F")
F_FIRMA = Font(name="Calibri", size=8, bold=True, color="475569")
AL_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)

SIN_DATO = "—"


def celda(ws, fila, col, valor=None, font=F_CELDA, fill=None, borde=True,
          formato=None, align=None):
    cel = ws.cell(row=fila, column=col)
    if valor is not None:
        cel.value = valor
    cel.font = font
    if fill:
        cel.fill = PatternFill("solid", fgColor=fill)
    if borde:
        cel.border = BORDE_TABLA
    cel.alignment = align or AL_CENTRO
    if formato:
        cel.number_format = formato
    return cel


def banda(ws, fila, c1, c2, texto, fill=SECCION):
    ws.merge_cells(start_row=fila, start_column=c1, end_row=fila, end_column=c2)
    celda(ws, fila, c1, texto, font=F_SECCION, fill=fill)
    for c in range(c1, c2 + 1):
        celda(ws, fila, c, fill=fill)


def par(ws, fila, col, etiqueta, valor, formato=None, ancho_val=1):
    celda(ws, fila, col, etiqueta, font=F_LAB, fill=LABEL_FILL, align=AL_IZQ)
    for k in range(1, ancho_val):
        celda(ws, fila, col + k, fill=LABEL_FILL)
    celda(ws, fila, col + ancho_val, fill=RESULT_FILL)
    v = ws.cell(row=fila, column=col + ancho_val)
    v.value = SIN_DATO if valor is None or valor == "" else valor
    v.font = F_VAL
    if formato and isinstance(valor, (int, float)):
        v.number_format = formato


def fmt(v, dec=2):
    if isinstance(v, (int, float)):
        return round(v, dec)
    return v


def construir(payload, ws):
    meta = payload.get("meta", {})
    r = payload.get("resultados", {})
    proctor = r.get("proctor") or {}
    cbr = r.get("cbr") or {}

    ws.title = "Ficha de Calicata"
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top = ws.page_margins.bottom = 0.45
    for letra, w in zip("ABCDEFGHIJ", [26, 13, 13, 13, 13, 13, 13, 13, 13, 10]):
        ws.column_dimensions[letra].width = w

    # ------------------------------------------------------------ cabecera
    ws.merge_cells("A1:I1")
    celda(ws, 1, 1, "GOBIERNO REGIONAL DEL CUSCO",
          font=Font(name="Calibri", size=11, bold=True, color=LABEL_TXT))
    ws.merge_cells("A2:I2")
    celda(ws, 2, 1, "Gerencia Regional de Transportes y Comunicaciones Cusco",
          font=Font(name="Calibri", size=9, color=LABEL_TXT))
    ws.merge_cells("A3:I3")
    celda(ws, 3, 1, "SUB GERENCIA DE COBERTURA Y COMUNICACIONES · UNIDAD FUNCIONAL DE ESTUDIOS Y PROYECTOS",
          font=Font(name="Calibri", size=8, color=LABEL_TXT))
    ws.merge_cells("A4:I4")
    celda(ws, 4, 1, "Laboratorio de Mecánica de Suelos, Materiales y Pavimentos",
          font=Font(name="Calibri", size=8, italic=True, color=LABEL_TXT))

    ws.merge_cells("A5:G5")
    celda(ws, 5, 1, "FICHA DE RESULTADOS DE CALICATA",
          font=Font(name="Calibri", size=14, bold=True, color=OSCURO))
    ws.merge_cells("H5:I5")
    celda(ws, 5, 8, "Revisión N° 0", font=Font(name="Calibri", size=8, color="64748B"))
    ws.merge_cells("A6:G6")
    celda(ws, 6, 1, meta.get("codigo", ""), font=Font(name="Calibri", size=11, bold=True),
          fill=RESULT_FILL)
    ws.merge_cells("H6:I6")
    celda(ws, 6, 8, f"Emitido: {meta.get('generado', '')}",
          font=Font(name="Calibri", size=8, color="64748B"))
    ws.row_dimensions[7].height = 6

    # ---------------------------------------------------------- ficha técnica
    banda(ws, 8, 1, 9, "DATOS GENERALES")
    par(ws, 9, 1, "Proyecto / Tramo:", meta.get("tramo_nombre"), ancho_val=3)
    par(ws, 9, 5, "Ubicación:", meta.get("progresiva_nombre"), ancho_val=2)
    par(ws, 10, 1, "Calicata:", meta.get("progresiva"))
    par(ws, 10, 3, "Estrato:", meta.get("estrato"))
    par(ws, 10, 5, "Profundidad:", meta.get("profundidad"), ancho_val=2)
    par(ws, 10, 8, "Lado:", meta.get("lado"))
    par(ws, 11, 1, "Coord. Este:", fmt(meta.get("coordenada_este")), formato="0.00")
    par(ws, 11, 3, "Coord. Norte:", fmt(meta.get("coordenada_norte")), formato="0.00")
    par(ws, 11, 5, "Cota (msnm):", fmt(meta.get("elevacion")), ancho_val=2)
    par(ws, 11, 8, "Fecha:", meta.get("fecha"))
    for f in (9, 10, 11):
        ws.row_dimensions[f].height = 18
    ws.row_dimensions[12].height = 6

    # ---------------------------------------------------- resumen de resultados
    banda(ws, 13, 1, 9, "RESUMEN DE RESULTADOS")
    fila = 14

    def linea_res(etiqueta, valor, formato=None, destacado=False):
        nonlocal fila
        celda(ws, fila, 1, etiqueta, font=F_LAB, fill=LABEL_FILL, align=AL_IZQ)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=3)
        fill = RESULT_FILL if destacado else CALC_FILL
        font = Font(name="Calibri", size=10, bold=True, color="0F172A") if destacado else F_CALC
        celda(ws, fila, 2, fill=fill)
        celda(ws, fila, 3, fill=fill)
        v = ws.cell(row=fila, column=2)
        v.value = SIN_DATO if valor is None else valor
        v.font = font
        if formato and isinstance(valor, (int, float)):
            v.number_format = formato
        fila += 1

    linea_res("Humedad Natural (%)", fmt(r.get("humedad_natural")))
    linea_res("Límite Líquido - L.L. (%)", fmt(r.get("limite_liquido")))
    linea_res("Límite Plástico - L.P. (%)", fmt(r.get("limite_plastico")))
    linea_res("Índice de Plasticidad - I.P. (%)", fmt(r.get("indice_plasticidad")))
    linea_res("Clasificación SUCS", r.get("sucs"), destacado=True)
    linea_res("Clasificación AASHTO (con I.G.)", r.get("aashto"), destacado=True)
    grava, arena, finos = r.get("grava"), r.get("arena"), r.get("finos")
    granulometria = (None if grava is None and finos is None
                     else f"{fmt(grava)} / {fmt(arena)} / {fmt(finos)}")
    linea_res("% Grava / % Arena / % Finos", granulometria)
    cu, cc = r.get("cu"), r.get("cc")
    linea_res("Cu / Cc", None if cu is None and cc is None else f"{fmt(cu)} / {fmt(cc)}")
    linea_res("Proctor: M.D.S. (g/cm³)", fmt(proctor.get("mds"), 3), formato="0.000")
    linea_res("Proctor: O.C.H. (%)", fmt(proctor.get("och")))
    linea_res("CBR de Diseño (%)", fmt(cbr.get("cbr_diseno")), destacado=True)
    linea_res("CBR: M.D.S. de los moldes (g/cm³)", fmt(cbr.get("mds_moldes"), 3), formato="0.000")
    linea_res("Expansión a 96 h (%)", fmt(cbr.get("expansion")))
    ws.row_dimensions[fila].height = 6

    # ------------------------------------------------ descripción geotécnica
    fila += 2
    banda(ws, fila, 1, 9, "DESCRIPCIÓN GEOTÉCNICA DEL ESTRATO")
    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 2, end_column=9)
    celda(ws, fila, 1, r.get("descripcion_geotecnica") or SIN_DATO,
          font=Font(name="Calibri", size=9, color="0F172A"), align=AL_IZQ)

    # ------------------------------------------------ advertencias
    advertencias = payload.get("advertencias") or []
    if advertencias:
        fila += 4
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
        celda(ws, fila, 1, "Observaciones del sistema:",
              font=Font(name="Calibri", size=8, italic=True, color="B45309"), align=AL_IZQ)
        for adv in advertencias:
            fila += 1
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
            celda(ws, fila, 1, f"• {adv}",
                  font=Font(name="Calibri", size=8, color="B45309"), align=AL_IZQ)

    # ------------------------------------------------ firmas
    fila += 5
    bloques = [("ESP. ENSAYOS GEOTÉCNICOS", 1, 3),
               ("ESP. SUELOS Y PAVIMENTOS", 4, 6),
               ("SUPERVISOR", 7, 9)]
    for _n, ci, cf in bloques:
        for c in range(ci, cf + 1):
            ws.cell(row=fila, column=c).border = Border(
                top=Side(style="thin", color="0F172A"))
    for texto, ci, cf in bloques:
        ws.merge_cells(start_row=fila + 1, start_column=ci, end_row=fila + 1, end_column=cf)
        celda(ws, fila + 1, ci, texto, font=F_FIRMA)

    ws.print_area = f"A1:I{fila + 2}"
    wb_calc = ws.parent
    wb_calc.calculation.fullCalcOnLoad = True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--payload", required=True)
    ap.add_argument("--salida", required=True)
    args = ap.parse_args()
    with open(args.payload, encoding="utf-8") as f:
        payload = json.load(f)
    wb = Workbook()
    construir(payload, wb.active)
    wb.save(args.salida)
    print("OK ficha generada")


if __name__ == "__main__":
    main()
