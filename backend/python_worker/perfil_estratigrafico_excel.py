# -*- coding: utf-8 -*-
"""
perfil_estratigrafico_excel.py
-------------------------------------------------------------------------------
Generador del reporte Excel del Perfil Estratigrafico (mecanica de suelos).

Replica visualmente la lamina de la pagina web (grilla de profundidad cada
0.10 m, columnas por progresiva con rellenos de material, fila de cota de
terreno y panel de resultados de ensayos), usando openpyxl.

Uso:
    python perfil_estratigrafico_excel.py --payload payload.json --salida out.xlsx

Todo el contenido viene del payload (JSON). Sin datos quemados: los colores y
patrones provienen de suelos_diccionario_nlp (patron_svg + color) via backend.
"""

import argparse
import json
import math
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Mapeo patron_svg (BD) -> patternType de Excel (estilo de rayado).
# Colores: fgColor = color de trama del material, bgColor = fondo aclarado.
# ---------------------------------------------------------------------------
PATRONES_EXCEL = {
    "grava": "gray125",
    "grava_arena": "lightTrellis",
    "arena": "gray0625",
    "limo": "lightHorizontal",
    "arcilla": "lightUp",
    "roca": "lightDown",
    "roca_fracturada": "darkGrid",
    "relleno": "lightVertical",
    "afirmado": "darkHorizontal",
    "generico": "lightGray",
}

PASO = 0.10  # metros por fila de la grilla de profundidad (igual que la web)

# Paleta espejo del CSS de la web
COLOR_GRILLA_MAYOR = "4A7EBB"
COLOR_GRILLA_MENOR = "C9D8E8"
COLOR_FONDO = "FFFFFF"
COLOR_HUECO = "F1F2F3"
COLOR_LABEL_FONDO = "F5F7F9"
COLOR_LABEL_TEXTO = "1C3D5F"
COLOR_COTA_FONDO = "EEF6E4"
COLOR_COTA_LABEL = "DCEAD2"
COLOR_NSAYOS = "FDF3C4"
COLOR_TIPO = "2C5F8A"
COLOR_GRUPO = "DBE9F5"
COLOR_GRUPO_FONDO = "F4F8FC"
COLOR_CELDA_ENSAYO = "FFFDE9"
COLOR_LABEL_ENSAYO = "CDE6C2"

# ---------------------------------------------------------------------------
# Utilidades de color
# ---------------------------------------------------------------------------

def color_hex(hex_str, fallback="999999"):
    """Normaliza '#RRGGBB'/'RRGGBB' -> 'RRGGBB'; fallback si es invalido."""
    if not isinstance(hex_str, str):
        return fallback
    h = hex_str.strip().lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    if len(h) >= 6 and all(c in "0123456789abcdefABCDEF" for c in h[:6]):
        return h[:6].upper()
    return fallback


def aclarar(hex_str, factor=0.88):
    r, g, b = (int(color_hex(hex_str)[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: max(0, min(255, round(c + (255 - c) * factor)))
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def relleno(patron_svg, color):
    """PatternFill de Excel para un material (trama + fondo aclarado)."""
    c = color_hex(color)
    pattern = PATRONES_EXCEL.get(patron_svg, "lightGray")
    return PatternFill(patternType=pattern, fgColor=c, bgColor=aclarar(c))


def relleno_solido(hex_str):
    return PatternFill(patternType="solid", fgColor=color_hex(hex_str))


# ---------------------------------------------------------------------------
# Geometria (mismo criterio que el frontend)
# ---------------------------------------------------------------------------

def profundidad_maxima(progresivas):
    m = 0.0
    for p in progresivas:
        for e in p.get("estratos") or []:
            try:
                f = float(e.get("profundidad_final") or 0)
                if f > m:
                    m = f
            except (TypeError, ValueError):
                continue
    return max(1.5, math.ceil(m * 10) / 10)


def paso_etiquetas(max_depth):
    if max_depth > 8:
        return 1.0
    if max_depth > 4:
        return 0.5
    if max_depth > 2:
        return 0.2
    return 0.1


def segmentos_progresiva(progresiva, max_depth):
    """Bloques (tipo, desde, hasta, estrato) igual que el frontend:
    estratos ordenados + huecos rellenados con generico."""
    estratos = []
    for e in progresiva.get("estratos") or []:
        try:
            ini = float(e.get("profundidad_inicial") or 0)
            fin = float(e.get("profundidad_final") or 0)
        except (TypeError, ValueError):
            continue
        if fin > ini:
            estratos.append((ini, fin, e))
    estratos.sort(key=lambda t: t[0])

    segs = []
    cursor = 0.0
    for ini, fin, e in estratos:
        fin = min(fin, max_depth)
        if ini > cursor + 0.001:
            segs.append({"tipo": "hueco", "desde": cursor, "hasta": ini})
        if fin > cursor:
            segs.append({"tipo": "estrato", "desde": max(ini, cursor),
                         "hasta": fin, "estrato": e})
            cursor = fin
        if cursor >= max_depth:
            break
    if cursor < max_depth:
        segs.append({"tipo": "hueco", "desde": cursor, "hasta": max_depth})
    return segs


def fmt_num(v, dec=2):
    try:
        n = float(v)
    except (TypeError, ValueError):
        return ""
    if math.isclose(n, round(n, dec)):
        return round(n, dec)
    return round(n, dec)


# ---------------------------------------------------------------------------
# Hoja 1: lamina del perfil
# ---------------------------------------------------------------------------

def dibujar_perfil(ws, payload):
    progresivas = payload.get("progresivas") or []
    panel = payload.get("panel") or {}
    filas_panel = panel.get("filas") or []
    valores = panel.get("valores") or {}
    n_ensayos = panel.get("nEnsayos") or {}

    tramo = payload.get("tramo") or {}
    n = len(progresivas)
    max_depth = profundidad_maxima(progresivas)
    n_filas_prof = int(round(max_depth / PASO))
    paso_lbl = paso_etiquetas(max_depth)

    col_labels = 1                      # A: etiquetas + eje izquierdo
    col_ini = 2                         # B: primera progresiva
    col_fin = col_ini + n - 1           # ultima progresiva
    col_eje_der = col_fin + 1           # eje derecho

    FILA_TITULO = 1
    FILA_NUM = 2
    FILA_PROG = 3
    FILA_PROF_INI = 4

    # --- Dimensiones -------------------------------------------------------
    ws.column_dimensions["A"].width = 20
    for c in range(col_ini, col_fin + 1):
        ws.column_dimensions[get_column_letter(c)].width = 4.6
    ws.column_dimensions[get_column_letter(col_eje_der)].width = 7

    ws.row_dimensions[FILA_TITULO].height = 22
    ws.row_dimensions[FILA_NUM].height = 12
    ws.row_dimensions[FILA_PROG].height = 58
    for r in range(FILA_PROF_INI, FILA_PROF_INI + n_filas_prof):
        ws.row_dimensions[r].height = 11

    # --- Titulo ------------------------------------------------------------
    ws.merge_cells(start_row=FILA_TITULO, start_column=1,
                   end_row=FILA_TITULO, end_column=col_eje_der)
    c = ws.cell(row=FILA_TITULO, column=1,
                value=f"PERFIL ESTRATIGRÁFICO — Tramo {tramo.get('codigo', '')}"
                      f"{(' — ' + tramo['nombre']) if tramo.get('nombre') else ''}")
    c.font = Font(bold=True, size=13, color=COLOR_LABEL_TEXTO)
    c.alignment = Alignment(horizontal="center", vertical="center")

    # --- Fila N° -----------------------------------------------------------
    ws.cell(row=FILA_NUM, column=col_labels, value="N°").font = Font(
        bold=True, size=7.5, color=COLOR_LABEL_TEXTO)
    for i, p in enumerate(progresivas):
        c = ws.cell(row=FILA_NUM, column=col_ini + i, value=i + 1)
        c.font = Font(bold=True, size=7.5)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = relleno_solido(COLOR_FONDO)

    # --- Fila Progresiva (km), codigo rotado como en la web ----------------
    ws.cell(row=FILA_PROG, column=col_labels, value="Progresiva (km)").font = Font(
        bold=True, size=7.5, color=COLOR_LABEL_TEXTO)
    for i, p in enumerate(progresivas):
        c = ws.cell(row=FILA_PROG, column=col_ini + i,
                    value=p.get("nombre") or p.get("codigo") or "")
        c.font = Font(size=6.5, bold=True, color=COLOR_LABEL_TEXTO)
        c.alignment = Alignment(horizontal="center", vertical="bottom",
                                textRotation=90, wrap_text=False)
        c.fill = relleno_solido("FDFBEA")

    # --- Grilla de profundidad + estratos -----------------------------------
    borde_menor = Side(style="thin", color=COLOR_GRILLA_MENOR)
    borde_mayor = Side(style="thin", color=COLOR_GRILLA_MAYOR)

    filas_label = set()
    for k in range(0, n_filas_prof + 1):
        d = round(k * PASO, 2)
        if math.isclose((d / paso_lbl) - round(d / paso_lbl), 0, abs_tol=1e-9):
            filas_label.add(k)

    for k in range(n_filas_prof):
        r = FILA_PROF_INI + k
        d_ini = round(k * PASO, 2)
        d_fin = round((k + 1) * PASO, 2)
        es_label = (k + 1) in filas_label
        lado = borde_mayor if es_label else borde_menor

        # Etiquetas de profundidad (izquierda y derecha, como la web)
        if k in filas_label:
            cl = ws.cell(row=r, column=col_labels, value=round(d_ini, 2))
            cl.font = Font(size=6.5)
            cl.alignment = Alignment(horizontal="right", vertical="top")
            cl.number_format = "0.00"
            cr = ws.cell(row=r, column=col_eje_der, value=round(d_ini, 2))
            cr.font = Font(size=6.5)
            cr.alignment = Alignment(horizontal="left", vertical="top")
            cr.number_format = "0.00"

        for i, p in enumerate(progresivas):
            col = col_ini + i
            celda = ws.cell(row=r, column=col)
            celda.border = Border(bottom=lado, left=borde_mayor)
            # contenido del estrato se pinta abajo (relleno por segmento)

    # borde derecho del bloque
    for k in range(n_filas_prof):
        r = FILA_PROF_INI + k
        es_label = (k + 1) in filas_label
        ws.cell(row=r, column=col_fin).border = Border(
            bottom=(borde_mayor if es_label else borde_menor),
            left=borde_mayor, right=borde_mayor)

    # Relleno de segmentos (estratos y huecos) por columna
    for i, p in enumerate(progresivas):
        col = col_ini + i
        for seg in segmentos_progresiva(p, max_depth):
            r_desde = FILA_PROF_INI + int(round(seg["desde"] / PASO))
            r_hasta = FILA_PROF_INI + int(round(seg["hasta"] / PASO))
            r_hasta = max(r_hasta, r_desde + 1)
            if seg["tipo"] == "hueco":
                fill = relleno_solido(COLOR_HUECO)
                texto = None
            else:
                e = seg["estrato"]
                patron = e.get("patron") or {}
                fill = relleno(patron.get("patron_svg") or "generico",
                               patron.get("color_hex_sugerido"))
                texto = (patron.get("clasificacion_sucs")
                         or (e.get("nombre") or ""))[:10] or None
            for r in range(r_desde, min(r_hasta, FILA_PROF_INI + n_filas_prof)):
                celda = ws.cell(row=r, column=col)
                celda.fill = fill
                if texto and r == (r_desde + r_hasta - 1) // 2:
                    celda.value = texto
                    celda.font = Font(size=5.5, color="22313F", bold=True)
                    celda.alignment = Alignment(horizontal="center",
                                                vertical="center")

    # --- Cota de terreno ------------------------------------------------------
    r_cota = FILA_PROF_INI + n_filas_prof
    ws.row_dimensions[r_cota].height = 13
    cl = ws.cell(row=r_cota, column=col_labels, value="Cota terreno (msnm)")
    cl.font = Font(bold=True, size=7.5, color="2F5233")
    cl.fill = relleno_solido(COLOR_COTA_LABEL)
    for i, p in enumerate(progresivas):
        celda = ws.cell(row=r_cota, column=col_ini + i,
                        value=fmt_num(p.get("elevacion")))
        celda.fill = relleno_solido(COLOR_COTA_FONDO)
        celda.font = Font(size=6.5, bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    # --- Panel de ensayos -------------------------------------------------------
    r = r_cota + 1
    fila_font = Font(size=7)

    def fila_encabezado(titulo, color_label, color_fondo, texto_blanco=False):
        nonlocal r
        ws.row_dimensions[r].height = 13
        cl = ws.cell(row=r, column=col_labels, value=titulo)
        cl.font = Font(bold=True, size=7.5,
                       color="FFFFFF" if texto_blanco else COLOR_LABEL_TEXTO)
        cl.fill = relleno_solido(color_label)
        for i in range(n):
            ws.cell(row=r, column=col_ini + i).fill = relleno_solido(color_fondo)
        r += 1

    def fila_datos(etiqueta, getter, fondo, negrita=False):
        nonlocal r
        ws.row_dimensions[r].height = 12
        cl = ws.cell(row=r, column=col_labels, value=etiqueta)
        cl.font = Font(bold=True, size=7, color="2F5233")
        cl.fill = relleno_solido(COLOR_LABEL_ENSAYO)
        cl.alignment = Alignment(vertical="center")
        for i, p in enumerate(progresivas):
            pid = str(p.get("id"))
            celda = ws.cell(row=r, column=col_ini + i, value=getter(p, pid))
            celda.fill = relleno_solido(fondo)
            celda.font = Font(size=6.5, bold=negrita)
            celda.alignment = Alignment(horizontal="center", vertical="center")
        r += 1

    # Fila N° de ensayos (etiqueta verde, celdas ámbar como en la web)
    ws.row_dimensions[r].height = 13
    cl = ws.cell(row=r, column=col_labels, value="Ensayos realizados (N°)")
    cl.font = Font(bold=True, size=7, color="2F5233")
    cl.fill = relleno_solido(COLOR_LABEL_ENSAYO)
    for i, p in enumerate(progresivas):
        pid = str(p.get("id"))
        celda = ws.cell(row=r, column=col_ini + i,
                        value=n_ensayos.get(pid, n_ensayos.get(p.get("id"), "")))
        celda.fill = relleno_solido(COLOR_NSAYOS)
        celda.font = Font(size=6.5, bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    r += 1

    for fila in filas_panel:
        if fila.get("mostrarTipo"):
            fila_encabezado(fila.get("tipoDescripcion", ""), COLOR_TIPO,
                            "EAF2FA", texto_blanco=True)
        if fila.get("mostrarGrupo") and fila.get("groupTitle"):
            fila_encabezado(fila["groupTitle"], COLOR_GRUPO, COLOR_GRUPO_FONDO)
        def getter(p, pid, _fila=fila):
            lista = valores.get(pid) or []
            return lista[_fila["_idx"]] if _fila.get("_idx") is not None and _fila["_idx"] < len(lista) else ""
        fila_datos(fila.get("label", ""), getter, COLOR_CELDA_ENSAYO)

    ws.freeze_panes = ws.cell(row=FILA_PROF_INI, column=col_ini)


# ---------------------------------------------------------------------------
# Hoja 2: detalle de estratos (editable)
# ---------------------------------------------------------------------------

def hoja_detalle(wb, payload):
    ws = wb.create_sheet("Detalle de Estratos")
    encabezados = ["ID (interno)", "N°", "Progresiva (km)", "Código", "Material",
                   "Descripción", "Prof. Inicial (m)", "Prof. Final (m)",
                   "Espesor (m)", "Clasif. SUCS", "Clasif. AASHTO",
                   "Patrón", "Cota terreno (msnm)", "N° Ensayos"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, size=9, color="FFFFFF")
        celda.fill = PatternFill(patternType="solid", fgColor=COLOR_TIPO)
        celda.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.column_dimensions["A"].hidden = True

    anchos = [10, 5, 14, 12, 22, 40, 12, 12, 11, 11, 12, 14, 13, 9]
    for idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = ancho

    fila = 2
    numero = 0
    for p in payload.get("progresivas") or []:
        for e in p.get("estratos") or []:
            numero += 1
            patron = e.get("patron") or {}
            try:
                ini = float(e.get("profundidad_inicial") or 0)
                fin = float(e.get("profundidad_final") or 0)
            except (TypeError, ValueError):
                ini, fin = None, None
            ws.cell(row=fila, column=1, value=e.get("id"))
            ws.cell(row=fila, column=2, value=numero)
            ws.cell(row=fila, column=3, value=p.get("nombre") or p.get("codigo"))
            ws.cell(row=fila, column=4, value=p.get("codigo"))
            ws.cell(row=fila, column=5, value=e.get("nombre"))
            ws.cell(row=fila, column=6, value=e.get("descripcion"))
            ws.cell(row=fila, column=7, value=ini).number_format = "0.00"
            ws.cell(row=fila, column=8, value=fin).number_format = "0.00"
            celda_e = ws.cell(row=fila, column=9)
            if ini is not None and fin is not None:
                celda_e.value = f"=IF(H{fila}=\"\",\"\",H{fila}-G{fila})"
            celda_e.number_format = "0.00"
            ws.cell(row=fila, column=10, value=e.get("nlp_clasificacion_sucs")
                    or patron.get("clasificacion_sucs"))
            ws.cell(row=fila, column=11, value=e.get("nlp_clasificacion_aashto"))
            ws.cell(row=fila, column=12, value=patron.get("nombre_material"))
            ws.cell(row=fila, column=13, value=fmt_num(p.get("elevacion")))
            ws.cell(row=fila, column=14, value=e.get("n_ensayos") or None)
            fila += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{max(fila - 1, 1)}"
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------------------
# Hoja 3: leyenda y metadatos
# ---------------------------------------------------------------------------

def hoja_leyenda(wb, payload):
    ws = wb.create_sheet("Leyenda y Metadatos")
    tramo = payload.get("tramo") or {}
    ws.append(["LEYENDA DE MATERIALES"])
    ws["A1"].font = Font(bold=True, size=12, color=COLOR_LABEL_TEXTO)

    ws.append(["Muestra", "Material", "Clasif. SUCS", "Clasif. AASHTO", "Patrón"])
    for celda in ws[2]:
        celda.font = Font(bold=True, size=9, color="FFFFFF")
        celda.fill = PatternFill(patternType="solid", fgColor=COLOR_TIPO)

    fila = 3
    vistos = set()
    for patron in payload.get("patrones") or []:
        clave = patron.get("nombre_original_excel")
        if clave in vistos:
            continue
        vistos.add(clave)
        muestra = ws.cell(row=fila, column=1)
        muestra.fill = relleno(patron.get("patron_svg"),
                               patron.get("color_hex_sugerido"))
        ws.cell(row=fila, column=2, value=clave)
        ws.cell(row=fila, column=3, value=patron.get("clasificacion_sucs"))
        ws.cell(row=fila, column=4, value=patron.get("clasificacion_aashto"))
        ws.cell(row=fila, column=5, value=patron.get("patron_svg"))
        fila += 1

    fila += 1
    meta = [
        ("Tramo", f"{tramo.get('codigo', '')} {tramo.get('nombre') or ''}".strip()),
        ("Modo", "Solo progresivas con ensayos" if payload.get("soloConDatos") else "Todas las progresivas"),
        ("Total progresivas", payload.get("totalProgresivas", "")),
        ("Fecha de exportación", payload.get("fecha", "")),
    ]
    for k, v in meta:
        ws.cell(row=fila, column=1, value=k).font = Font(bold=True, size=9)
        ws.cell(row=fila, column=2, value=v).font = Font(size=9)
        fila += 1

    for col, ancho in zip("ABCDE", (16, 34, 12, 13, 16)):
        ws.column_dimensions[col].width = ancho
    return ws


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--payload", required=True)
    parser.add_argument("--salida", required=True)
    args = parser.parse_args()

    with open(args.payload, "r", encoding="utf-8") as f:
        payload = json.load(f)

    # Indices de filas del panel para acceso posicional a valores
    filas = payload.get("panel", {}).get("filas") or []
    for idx, fila in enumerate(filas):
        fila["_idx"] = idx

    wb = Workbook()
    ws = wb.active
    ws.title = "Perfil Estratigráfico"
    dibujar_perfil(ws, payload)
    hoja_detalle(wb, payload)
    hoja_leyenda(wb, payload)

    wb.save(args.salida)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"error": str(exc)}), file=sys.stderr)
        sys.exit(1)
